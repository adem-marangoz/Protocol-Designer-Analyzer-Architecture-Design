"""Export a protocol to a formatted Excel workbook (.xlsx).

The workbook is meant to be read by people, so it favours clarity over raw
data: a title page with tiles and clickable contents, one filterable table of
messages, colour-coded byte maps of every frame, one sheet per message, a flat
table of all fields, and sheets for enumerations, CRCs, tests, simulator rules
and validation results. Colours match the PDF specification.
"""

from __future__ import annotations

import datetime as _dt
import json
import re
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple, Union

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from .. import APP_NAME, __version__
from ..protocol.decoder import Decoder
from ..protocol.encoder import Encoder
from ..protocol.errors import DefinitionError, ProtocolError
from ..protocol.fields import crc_algorithm, field_endianness, resolve_enum
from ..protocol.layout import covering_length_field, crc_range, length_range
from ..protocol.model import Encoding, FieldType, FrameDefinition, ProtocolDefinition
from ..protocol.validator import ERROR, validate_protocol
from ..storage.json_writer import step_to_dict
from .spec_report import frame_offsets, frame_size_text, range_text, rule_text, type_text

# ------------------------------------------------------------------ style ---

NAVY = "1F3A5F"
WHITE = "FFFFFF"
MUTED = "57606A"
KEY_BG = "EEF2F7"
ZEBRA = "F6F8FA"
GRID = "C9D1D9"
LINK = "0B57D0"

ENCODING_FILL = {
    Encoding.CONSTANT: "DBEAFE",
    Encoding.LENGTH: "FEF3C7",
    Encoding.CRC: "FDE2E1",
    Encoding.VALUE: "E6F4EA",
}
ENCODING_INK = {
    Encoding.CONSTANT: "1D4ED8",
    Encoding.LENGTH: "B45309",
    Encoding.CRC: "B91C1C",
    Encoding.VALUE: "15803D",
}
ENCODING_LABEL = {
    Encoding.CONSTANT: "Constant",
    Encoding.LENGTH: "Length",
    Encoding.CRC: "CRC",
    Encoding.VALUE: "Value",
}
ENCODING_MEANING = {
    Encoding.CONSTANT: "Constant: always the same value; used to recognise the message (e.g. start byte, command).",
    Encoding.LENGTH: "Length: calculated automatically from the size of the payload.",
    Encoding.CRC: "Checksum / CRC: calculated automatically over the bytes shown in its rule.",
    Encoding.VALUE: "Data: the value you send or receive.",
}
DIRECTION_FILL = {"TX": "DBEAFE", "RX": "E6F4EA", "BOTH": "EDEFF2"}
DIRECTION_TAB = {"TX": "1D4ED8", "RX": "15803D", "BOTH": "6E7781"}

FONT = "Calibri"
MONO = "Consolas"


def font(**kw) -> Font:
    kw.setdefault("name", FONT)
    return Font(**kw)


THIN = Side(style="thin", color=GRID)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = font(bold=True, color=WHITE)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

BYTES_PER_LINE = 16
FIELD_HEADERS = ["#", "Field", "Byte offset", "Size (bytes)", "Type", "Byte order", "Encoding", "Rule / value",
                 "Scale", "Offset (value)", "Unit", "Min", "Max", "Enumeration", "Default", "Description"]
FIELD_WIDTHS = [5, 20, 11, 11, 14, 12, 12, 38, 8, 12, 8, 8, 8, 18, 10, 36]
# message sheets also show the example packet next to each field
MESSAGE_HEADERS = FIELD_HEADERS[:-1] + ["Example (hex)", "Example value", "Description"]
MESSAGE_WIDTHS = FIELD_WIDTHS[:-1] + [16, 24, 36]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


class Link:
    """A cell value that jumps to another sheet when clicked."""

    def __init__(self, text: str, sheet: str, cell: str = "A1"):
        self.text = text
        self.location = "'" + sheet.replace("'", "''") + f"'!{cell}"

    def hyperlink(self) -> Hyperlink:
        return Hyperlink(ref="", location=self.location)


def _plain(value: Any) -> Any:
    if isinstance(value, (dict, list, bool)) or value is None:
        return json.dumps(value, ensure_ascii=False) if value is not None else ""
    return value


def _value_text(value: Any) -> str:
    """A field value as people write it; bit-field dicts become ``(name = value, ...)``."""
    if isinstance(value, dict):
        return "(" + ", ".join(f"{k} = {_plain(v)}" for k, v in value.items()) + ")"
    return str(_plain(value))


def _hex(value: Optional[int], width: int = 2) -> str:
    return "" if value is None else f"0x{value:0{width}X}"


def _step_details(d: Dict[str, Any]) -> str:
    """Test step parameters as text: ``values`` become ``NAME = value`` pairs."""
    d = dict(d)
    parts = []
    if "variable" in d:
        parts.append(f"{d.pop('variable')} = {_plain(d.pop('value', ''))}")
    for key, value in d.items():
        if key == "values" and isinstance(value, dict):
            parts += [f"{k} = {_value_text(v)}" for k, v in value.items()]
        elif key == "delay_ms":
            parts.append(f"wait {value} ms")
        elif key == "data":
            parts.append(f"bytes {value}")
        elif key == "text":
            parts.append(str(value))
        else:
            parts.append(f"{key} = {_plain(value)}")
    return "\n".join(parts)


_CHECK_WORDS = {"equals": "=", "not_equals": "≠", "min": "≥", "max": "≤"}


def _check_text(check: Dict[str, Any]) -> str:
    """A test check in words, e.g. ``PRESSURE between 1.5 and 3.5``."""
    check = dict(check)
    words = []
    if "min" in check and "max" in check:
        words.append(f"between {_plain(check.pop('min'))} and {_plain(check.pop('max'))}")
    for key, value in check.items():
        if key in _CHECK_WORDS:
            words.append(f"{_CHECK_WORDS[key]} {_plain(value)}")
        elif key == "save_as":
            words.append(f"→ save as ${{{value}}}")
        elif key == "raw" and value:
            words.append("(raw value)")
        elif key != "field":
            words.append(f"{key} {_plain(value)}")
    return f"{check.get('field', '')} " + " ".join(words)


# --------------------------------------------------------------- writers ---


def set_widths(ws: Worksheet, widths: Sequence[float], first_col: int = 1) -> None:
    for i, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(first_col + i)].width = width


def write(ws: Worksheet, row: int, col: int, value: Any, **style) -> Any:
    cell = ws.cell(row=row, column=col)
    if isinstance(value, Link):
        cell.value = value.text
        cell.hyperlink = value.hyperlink()
        cell.font = font(color=LINK, underline="single", bold=style.pop("bold", False))
    else:
        value = _plain(value)
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("\ufffd", value)
        cell.value = value
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"  # text, not a formula
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def banner(ws: Worksheet, title: str, subtitle: str, width_cols: int, back: Optional[Link] = None) -> int:
    """Navy title band across the top of a sheet. Returns the next free row."""
    ws.sheet_view.showGridLines = False
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    write(ws, 1, 1, title, font=font(bold=True, size=18, color=WHITE), fill=HEADER_FILL,
          alignment=Alignment(vertical="center", indent=1))
    for c in range(2, width_cols + 1):
        ws.cell(row=1, column=c).fill = HEADER_FILL
    ws.row_dimensions[1].height = 36
    ws.merge_cells(f"A2:{last}2")
    write(ws, 2, 1, subtitle, font=font(italic=True, color=MUTED), alignment=Alignment(vertical="center", indent=1))
    ws.row_dimensions[2].height = 20
    if back is not None:
        write(ws, 3, 1, back)
    return 4


def section(ws: Worksheet, row: int, text: str, width_cols: int) -> int:
    write(ws, row, 1, text, font=font(bold=True, size=13, color=NAVY))
    for c in range(1, width_cols + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[row].height = 22
    return row + 1


def key_values(ws: Worksheet, row: int, pairs: List[Tuple[Any, ...]], value_span: int = 1, col: int = 1,
               chars: int = 60) -> int:
    """Label / value rows. A pair may carry a third item: extra style for the value.

    The value is merged across ``value_span`` columns; Excel cannot size merged
    rows itself, so long values get a height from ``chars`` (characters per line).
    """
    for key, value, *extra in pairs:
        write(ws, row, col, key, font=font(bold=True), fill=fill(KEY_BG), border=BOX, alignment=WRAP)
        write(ws, row, col + 1, value, border=BOX, alignment=WRAP, **(extra[0] if extra else {}))
        if value_span > 1:
            ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + value_span)
            for c in range(col + 2, col + 1 + value_span):
                ws.cell(row=row, column=c).border = BOX
        if isinstance(value, str) and len(value) > chars:
            ws.row_dimensions[row].height = 15 * (1 + len(value) // chars)
        row += 1
    return row + 1


CellStyle = Callable[[int, int, Any], Dict[str, Any]]


def table(
    ws: Worksheet,
    row: int,
    headers: List[str],
    rows: List[List[Any]],
    style: Optional[CellStyle] = None,
    band: Optional[List[int]] = None,
    filter_: bool = False,
    freeze: bool = False,
) -> int:
    """Header + body rows with borders and zebra banding. Returns the next free row.

    ``style(row_index, col_index, value)`` may return extra cell attributes;
    ``band`` gives a band number per row (rows of the same group share a colour).
    """
    for c, h in enumerate(headers, 1):
        write(ws, row, c, h, font=HEADER_FONT, fill=HEADER_FILL, border=BOX, alignment=CENTER)
    ws.row_dimensions[row].height = 30
    header_row = row
    for r, values in enumerate(rows):
        row += 1
        shade = (band[r] if band else r) % 2 == 1
        for c, value in enumerate(values, 1):
            attrs: Dict[str, Any] = {"border": BOX, "alignment": WRAP}
            if shade:
                attrs["fill"] = fill(ZEBRA)
            if style:
                attrs.update(style(r, c, value) or {})
            write(ws, row, c, value, **attrs)
    if filter_ and rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{row}"
    if freeze:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return row + 2


def print_setup(ws: Worksheet, title: str) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = title
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_options.horizontalCentered = True


# ------------------------------------------------------------ sheet names ---

_INVALID = re.compile(r"[\[\]:*?/\\]")


class SheetNames:
    """Excel sheet names: max 31 characters, no []:*?/\\, unique ignoring case, not "History"."""

    def __init__(self, reserved: Sequence[str]):
        self.used = {n.lower() for n in reserved} | {"history"}

    def make(self, wanted: str) -> str:
        base = _INVALID.sub("_", wanted).strip("'").strip() or "Sheet"
        base = base[:31]
        name, n = base, 2
        while name.lower() in self.used:
            suffix = f" ({n})"
            name = base[: 31 - len(suffix)] + suffix
            n += 1
        self.used.add(name.lower())
        return name


FIXED_SHEETS = ["Overview", "Messages", "Byte maps", "All fields", "Enumerations", "CRC", "Tests", "Simulator",
                "Validation"]


# --------------------------------------------------------------- analysis ---


def example_packet(protocol: ProtocolDefinition, frame: FrameDefinition):
    """(encoded, decoded) with default values, or (None, None) if the frame cannot be built."""
    try:
        encoded = Encoder(protocol).encode(frame, {})
        decoded = Decoder(protocol).decode(frame, encoded.data, check_constants=False)
        return encoded, decoded
    except (ProtocolError, KeyError, ValueError):
        return None, None


def field_row(protocol: ProtocolDefinition, frame: FrameDefinition, index: int, offset, size) -> List[Any]:
    f = frame.fields[index]
    try:
        type_label = type_text(f, size, protocol)
    except DefinitionError:
        type_label = f.type.value
    multi = (size or 0) > 1 and f.type not in (FieldType.BYTES, FieldType.STRING)
    order = ""
    if multi:
        order = "little endian" if field_endianness(f, protocol).value == "LITTLE" else "big endian"
    numeric = f.encoding == Encoding.VALUE and f.type not in (FieldType.BYTES, FieldType.STRING, FieldType.BOOLEAN,
                                                              FieldType.BITFIELD)
    enum = ""
    if f.enum is not None:
        enum = f.enum if isinstance(f.enum, str) else ", ".join(f"{k}={v}" for k, v in sorted(f.enum.items()))
    return [
        index + 1,
        f.name,
        "" if offset is None else offset,
        "variable" if size is None else size,
        type_label,
        order,
        ENCODING_LABEL[f.encoding],
        rule_text(frame, index, f),
        f.scale if numeric and f.is_scaled else "",
        f.value_offset if numeric and f.value_offset else "",
        f.unit,
        "" if f.minimum is None else f.minimum,
        "" if f.maximum is None else f.maximum,
        enum,
        "" if f.default is None else f.default,
        f.description,
    ]


def field_style(frame: FrameDefinition, name_col: int, encoding_col: int) -> CellStyle:
    def style(r: int, c: int, value: Any) -> Dict[str, Any]:
        enc = frame.fields[r].encoding
        if c == name_col:
            return {"fill": fill(ENCODING_FILL[enc]), "font": font(bold=True, color=ENCODING_INK[enc])}
        if c == encoding_col:
            return {"font": font(color=ENCODING_INK[enc])}
        if c in (1, 3, 4):
            return {"alignment": Alignment(horizontal="center", vertical="center")}
        return {}
    return style


# ----------------------------------------------------------------- sheets ---


class WorkbookBuilder:
    def __init__(self, protocol: ProtocolDefinition, source: Optional[str]):
        self.p = protocol
        self.source = source
        self.wb = Workbook()
        names = SheetNames(FIXED_SHEETS)
        self.frame_sheet = {f.name: names.make(f.name) for f in protocol.frames}
        self.bytemap_cell: Dict[str, str] = {}
        self.issues = validate_protocol(protocol)
        self.examples = {f.name: example_packet(protocol, f) for f in protocol.frames}

    def build(self) -> Workbook:
        overview = self.wb.active
        overview.title = "Overview"
        self.messages_sheet()
        self.bytemaps_sheet()
        for frame in self.p.frames:
            self.frame_sheet_(frame)
        self.all_fields_sheet()
        self.enums_sheet()
        self.crc_sheet()
        self.tests_sheet()
        self.simulator_sheet()
        self.validation_sheet()
        self.overview_sheet(overview)  # last: it links to everything above
        for ws in self.wb.worksheets:
            print_setup(ws, f"{self.p.name} v{self.p.version}")
        props = self.wb.properties
        props.title = f"{self.p.name} – Protocol"
        props.creator = f"{APP_NAME} {__version__}"
        props.description = self.p.description
        return self.wb

    # ------------------------------------------------------------ overview ---

    def overview_sheet(self, ws: Worksheet) -> None:
        p = self.p
        set_widths(ws, [24, 18, 18, 18, 18, 18, 18])
        row = banner(ws, p.name, f"Protocol specification · version {p.version} · "
                                 f"{_dt.datetime.now():%Y-%m-%d %H:%M} · {APP_NAME} {__version__}", 7)
        ws.sheet_properties.tabColor = NAVY

        # tiles
        n_fields = sum(len(f.fields) for f in p.frames)
        errors = sum(1 for i in self.issues if i.severity == ERROR)
        tiles = [("Messages", len(p.frames), "Messages"), ("Fields", n_fields, "All fields"),
                 ("Enumerations", len(p.enums), "Enumerations"), ("Tests", len(p.tests), "Tests"),
                 ("Problems", errors, "Validation")]
        for i, (label, number, target) in enumerate(tiles):
            col = 2 + i
            colour = "B91C1C" if label == "Problems" and number else NAVY
            num = ws.cell(row=row, column=col, value=number)
            num.font = font(bold=True, size=22, color=colour)
            num.alignment = CENTER
            num.fill = fill(KEY_BG)
            num.hyperlink = Link(label, target).hyperlink()
            lab = ws.cell(row=row + 1, column=col, value=label)
            lab.font = font(color=MUTED)
            lab.alignment = CENTER
            lab.fill = fill(KEY_BG)
            for cell in (num, lab):
                cell.border = Border(left=Side(style="thin", color=WHITE), right=Side(style="thin", color=WHITE))
        ws.row_dimensions[row].height = 38
        row += 3

        row = section(ws, row, "Protocol", 7)
        pairs: List[Tuple[str, Any]] = [
            ("Name", p.name), ("Version", p.version), ("Description", p.description or "—"),
            ("Default byte order", "Big endian (most significant byte first)" if p.endianness.value == "BIG"
             else "Little endian (least significant byte first)"),
        ]
        if self.source:
            pairs.append(("Source file", self.source))
        row = key_values(ws, row, pairs, value_span=6, chars=110)

        t = p.transport
        row = section(ws, row, "Transport", 7)
        row = key_values(ws, row, [
            ("Type", t.type.value), ("Summary", t.summary()),
            ("Serial", f"port {t.port or '—'}, {t.baudrate} baud, {t.data_bits} data bits, parity {t.parity}, "
                       f"{t.stop_bits} stop bit(s)"),
            ("Network", f"{t.host}:{t.tcp_port}"),
            ("CAN", f"{t.can_interface} channel {t.can_channel}, {t.can_bitrate} bit/s"),
            ("Timeout", f"{t.timeout_ms} ms"),
        ], value_span=6, chars=110)

        row = section(ws, row, "Contents (click to open)", 7)
        contents = [
            ("Messages", "All messages in one table: ID, direction, size, response."),
            ("Byte maps", "Picture of every message, byte by byte, with an example packet."),
        ]
        contents += [(self.frame_sheet[f.name], f"Message {f.name}: fields, bit fields, rules and example.")
                     for f in p.frames]
        contents += [
            ("All fields", "Every field of every message in one filterable table."),
            ("Enumerations", "Value → name tables."),
            ("CRC", "Checksum algorithms with all parameters and check values."),
            ("Tests", "Automated test steps."),
            ("Simulator", "How the built-in simulated device answers."),
            ("Validation", "Problems found in the protocol definition."),
        ]
        for sheet, text in contents:
            write(ws, row, 1, Link(sheet, sheet))
            write(ws, row, 2, text, font=font(color=MUTED))
            row += 1
        row += 1

        row = section(ws, row, "Colour legend", 7)
        for enc in (Encoding.CONSTANT, Encoding.LENGTH, Encoding.CRC, Encoding.VALUE):
            write(ws, row, 1, ENCODING_LABEL[enc], fill=fill(ENCODING_FILL[enc]),
                  font=font(bold=True, color=ENCODING_INK[enc]), border=BOX, alignment=CENTER)
            write(ws, row, 2, ENCODING_MEANING[enc])
            row += 1
        row += 1

        row = section(ws, row, "How to read this workbook", 7)
        for line in [
            "• Byte offsets count from the first byte of the message (offset 0).",
            "• Multi-byte numbers are sent in the byte order shown (big endian = most significant byte first).",
            "• 'raw' is the number on the wire; the real value = raw × Scale + Offset (value), in the given Unit.",
            "• Enumerations translate numbers into names (e.g. 0x01 = ACTIVE).",
            "• LENGTH and CRC fields are filled in automatically; their rule says what they cover.",
            "• Click a sheet name above (or a message name anywhere) to jump to it; '← Overview' leads back.",
            "• The Messages, All fields, Enumerations and Tests tables have filter buttons in their header row.",
        ]:
            write(ws, row, 1, line)
            row += 1
        ws.freeze_panes = None

    # ------------------------------------------------------------ messages ---

    def messages_sheet(self) -> None:
        ws = self.wb.create_sheet("Messages")
        ws.sheet_properties.tabColor = NAVY
        widths = [22, 9, 9, 10, 17, 18, 18, 8, 50]
        set_widths(ws, widths)
        row = banner(ws, "Messages", f"{len(self.p.frames)} message(s) · click a name to open its sheet",
                     len(widths), Link("← Overview", "Overview"))
        rows = []
        for f in self.p.frames:
            rows.append([
                Link(f.name, self.frame_sheet[f.name]),
                _hex(f.id), "" if f.id is None else f.id,
                f.direction.value,
                (_hex(f.can_id, 8 if f.can_extended else 3) + (" (ext)" if f.can_extended else ""))
                if f.can_id is not None else "",
                frame_size_text(f),
                f.expected_response or "",
                len(f.fields),
                f.description,
            ])

        def style(r, c, value):
            if c == 4:
                return {"fill": fill(DIRECTION_FILL.get(value, "FFFFFF")), "alignment": CENTER,
                        "font": font(bold=True, color=DIRECTION_TAB.get(value, MUTED))}
            if c in (2, 3, 5, 8):
                return {"alignment": Alignment(horizontal="center", vertical="center")}
            return {}

        table(ws, row, ["Message", "ID (hex)", "ID (dec)", "Direction", "CAN ID", "Size", "Expected response",
                        "Fields", "Description"], rows, style=style, filter_=True, freeze=True)

    # ----------------------------------------------------------- byte maps ---

    def bytemaps_sheet(self) -> None:
        ws = self.wb.create_sheet("Byte maps")
        ws.sheet_properties.tabColor = NAVY
        ws.column_dimensions["A"].width = 12
        for c in range(2, 2 + BYTES_PER_LINE):
            ws.column_dimensions[get_column_letter(c)].width = 10
        row = banner(ws, "Byte maps", "Each message byte by byte: byte number, field, and an example packet "
                                      "built with default values (hex).", 1 + BYTES_PER_LINE,
                     Link("← Overview", "Overview"))
        # legend
        for i, enc in enumerate((Encoding.CONSTANT, Encoding.LENGTH, Encoding.CRC, Encoding.VALUE)):
            col = 2 + i * 3
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            write(ws, row, col, ENCODING_LABEL[enc] + (" (data)" if enc == Encoding.VALUE else " (automatic)"
                                                       if enc in (Encoding.LENGTH, Encoding.CRC) else " (fixed)"),
                  fill=fill(ENCODING_FILL[enc]),
                  font=font(bold=True, color=ENCODING_INK[enc]), alignment=CENTER)
        write(ws, row, 1, "Legend", font=font(bold=True, color=MUTED))
        row += 2

        for frame in self.p.frames:
            self.bytemap_cell[frame.name] = f"A{row}"
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + BYTES_PER_LINE)
            title = f"{frame.name}   ·   {frame.direction.value}   ·   {frame_size_text(frame)}"
            cell = write(ws, row, 1, Link(title, self.frame_sheet[frame.name]), bold=True)
            cell.font = font(bold=True, size=12, color=LINK, underline="single")
            ws.row_dimensions[row].height = 22
            row = self._byte_grid(ws, row + 1, frame) + 1

    def _byte_cells(self, frame: FrameDefinition) -> List[Tuple[int, Optional[int], str]]:
        """(field index, byte offset or None, hex text) for every byte of the example packet."""
        encoded, _ = self.examples[frame.name]
        cells: List[Tuple[int, Optional[int], str]] = []
        if encoded is not None:
            for span in encoded.spans:
                if span.size == 0:
                    cells.append((span.index, None, "(empty)"))
                for k in range(span.size):
                    cells.append((span.index, span.offset + k, f"{encoded.data[span.offset + k]:02X}"))
            return cells
        for i, (f, offset, size) in enumerate(frame_offsets(frame)):
            for k in range(size or 1):
                cells.append((i, None if offset is None or size is None else offset + k, "?"))
        return cells

    def _byte_grid(self, ws: Worksheet, row: int, frame: FrameDefinition) -> int:
        cells = self._byte_cells(frame)
        for start in range(0, max(len(cells), 1), BYTES_PER_LINE):
            chunk = cells[start : start + BYTES_PER_LINE]
            for label_row, label in enumerate(("Byte", "Field", "Example")):
                write(ws, row + label_row, 1, label, font=font(bold=True, color=MUTED),
                      alignment=Alignment(horizontal="right", vertical="center"))
            # byte numbers + hex
            for k, (index, offset, hex_text) in enumerate(chunk):
                col = 2 + k
                enc = frame.fields[index].encoding
                write(ws, row, col, "var" if offset is None else offset, font=font(size=9, color=MUTED),
                      alignment=CENTER, border=BOX)
                write(ws, row + 2, col, hex_text, font=font(name=MONO, size=10, bold=True),
                      alignment=CENTER, border=BOX, fill=fill(ENCODING_FILL[enc]))
            # field names, merged across their bytes
            k = 0
            while k < len(chunk):
                index = chunk[k][0]
                end = k
                while end + 1 < len(chunk) and chunk[end + 1][0] == index:
                    end += 1
                f = frame.fields[index]
                continued = start + k > 0 and cells[start + k - 1][0] == index
                label = f.name + (" (cont.)" if continued else "")
                c1, c2 = 2 + k, 2 + end
                if c2 > c1:
                    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 1, end_column=c2)
                write(ws, row + 1, c1, label, fill=fill(ENCODING_FILL[f.encoding]),
                      font=font(bold=True, size=9, color=ENCODING_INK[f.encoding]),
                      alignment=Alignment(horizontal="center", vertical="center", wrap_text=c2 > c1,
                                          shrink_to_fit=c1 == c2))
                for c in range(c1, c2 + 1):
                    ws.cell(row=row + 1, column=c).border = BOX
                    ws.cell(row=row + 1, column=c).fill = fill(ENCODING_FILL[f.encoding])
                k = end + 1
            ws.row_dimensions[row + 1].height = 30
            ws.row_dimensions[row + 2].height = 18
            row += 3
        return row

    # ------------------------------------------------------- message sheets ---

    def frame_sheet_(self, frame: FrameDefinition) -> None:
        name = self.frame_sheet[frame.name]
        ws = self.wb.create_sheet(name)
        ws.sheet_properties.tabColor = DIRECTION_TAB.get(frame.direction.value, MUTED)
        set_widths(ws, MESSAGE_WIDTHS)
        width = len(MESSAGE_HEADERS)
        row = banner(ws, frame.name, frame.description or f"Message {frame.name}", width,
                     Link("← All messages", "Messages"))
        if frame.name in self.bytemap_cell:
            write(ws, 3, 3, Link("Byte map →", "Byte maps", self.bytemap_cell[frame.name]))
        encoded, decoded = self.examples[frame.name]

        row = section(ws, row, "Message", width)
        pairs: List[Tuple[Any, ...]] = [
            ("Message ID", f"{_hex(frame.id)} ({frame.id})" if frame.id is not None else "—"),
            ("Direction", {"TX": "TX – sent to the device", "RX": "RX – received from the device"}
             .get(frame.direction.value, "both directions")),
            ("Frame size", frame_size_text(frame)),
            ("Expected response", Link(frame.expected_response, self.frame_sheet[frame.expected_response])
             if frame.expected_response in self.frame_sheet else "—"),
        ]
        if frame.can_id is not None:
            pairs.append(("CAN identifier", f"{_hex(frame.can_id, 8 if frame.can_extended else 3)} "
                                            f"({'29-bit extended' if frame.can_extended else '11-bit standard'})"))
        callers = [f.name for f in self.p.frames if f.expected_response == frame.name]
        if callers:
            pairs.append(("Response to", ", ".join(callers)))
        if encoded is not None:
            pairs.append(("Example packet", encoded.hex, {"font": font(name=MONO, bold=True)}))
        else:
            pairs.append(("Example packet", "none: the message definition has errors (see the Validation sheet)",
                          {"font": font(color="B91C1C")}))
        row = key_values(ws, row, pairs, value_span=6, col=2, chars=100)

        row = section(ws, row, "Fields", width)
        if any(f.bits for f in frame.fields):
            write(ws, row, 2, "Rows marked ↳ are the bits inside the bit field above them.",
                  font=font(italic=True, color=MUTED))
            row += 1
        shown = {fd.name: fd for fd in decoded.fields} if decoded is not None else {}
        rows: List[List[Any]] = []
        kinds: List[Tuple[Encoding, bool]] = []
        bands: List[int] = []
        for i, (f, offset, size) in enumerate(frame_offsets(frame)):
            fd = shown.get(f.name)
            values = field_row(self.p, frame, i, offset, size)
            example = [fd.hex or "(empty)", fd.display] if fd is not None else ["", ""]
            rows.append(values[:-1] + example + values[-1:])
            kinds.append((f.encoding, False))
            bands.append(i)
            for b in sorted(f.bits, key=lambda x: x.start):
                rows.append(self._bit_row(f, b, offset, fd))
                kinds.append((f.encoding, True))
                bands.append(i)

        def style(r, c, value):
            enc, is_bit = kinds[r]
            if c == 2:
                if is_bit:
                    return {"font": font(italic=True, color=ENCODING_INK[enc]), "alignment": Alignment(
                        vertical="center", wrap_text=True, indent=1)}
                return {"fill": fill(ENCODING_FILL[enc]), "font": font(bold=True, color=ENCODING_INK[enc])}
            if c == 7:
                return {"font": font(color=ENCODING_INK[enc])}
            if c == 16:
                return {"font": font(name=MONO), "alignment": Alignment(horizontal="center", vertical="center")}
            if c in (1, 3, 4, 15):
                return {"alignment": Alignment(horizontal="center", vertical="center")}
            return {}

        row = table(ws, row, MESSAGE_HEADERS, rows, style=style, band=bands)

        rules = self._rules(frame)
        if rules:
            row = section(ws, row, "How the message is built", width)
            for line in rules:
                write(ws, row, 1, "•", alignment=Alignment(horizontal="right"))
                write(ws, row, 2, line)
                row += 1

    def _bit_row(self, f, b, offset, fd) -> List[Any]:
        """A row of the message field table describing one bit group of a BITFIELD."""
        bits = f"bit {b.start}" if b.length == 1 else f"bits {b.start}–{b.start + b.length - 1}"
        try:
            mapping = resolve_enum(b.enum, self.p) if b.enum is not None else None
        except DefinitionError:
            mapping = None
        meaning = ", ".join(f"{v} = {n}" for v, n in sorted(mapping.items())) if mapping else (
            "1 = yes / set, 0 = no / clear" if b.length == 1 else f"number 0 … {(1 << b.length) - 1}")
        shown = next((x.display for x in fd.bits if x.name == b.name), "") if fd is not None else ""
        return ["", f"↳ {b.name}", "" if offset is None else offset, bits,
                "flag" if b.length == 1 else f"{b.length} bits", "", "", f"mask 0x{b.mask:X}: {meaning}",
                "", "", "", "", "", b.enum if isinstance(b.enum, str) else "", "", "", shown, b.description]

    def _rules(self, frame: FrameDefinition) -> List[str]:
        rules = []
        for i, f in enumerate(frame.fields):
            try:
                if f.encoding == Encoding.LENGTH:
                    start, end = length_range(frame, i)
                    rules.append(f"{f.name} = number of bytes of {range_text(frame, start, end)}"
                                 + (f" {f.length_adjust:+d}" if f.length_adjust else "") + ".")
                elif f.encoding == Encoding.CRC:
                    start, end = crc_range(frame, i)
                    rules.append(f"{f.name} = {crc_algorithm(f).name} over the bytes of "
                                 f"{range_text(frame, start, end)} (parameters on the CRC sheet).")
            except (KeyError, DefinitionError):
                continue
        for i, f in enumerate(frame.fields):
            if f.is_variable:
                rules.append(f"{f.name} has a variable size, given by the LENGTH field." if
                             covering_length_field(frame, i) is not None else
                             f"{f.name} has a variable size: all bytes before the trailing fields.")
        constants = [f.name for f in frame.fields if f.encoding == Encoding.CONSTANT]
        if constants:
            rules.append(f"The message is recognised by its constant field(s): {', '.join(constants)}"
                         + (" and its CAN identifier." if frame.can_id is not None else "."))
        elif frame.can_id is not None:
            rules.append("The message is recognised by its CAN identifier.")
        return rules

    # ---------------------------------------------------------- all fields ---

    def all_fields_sheet(self) -> None:
        ws = self.wb.create_sheet("All fields")
        ws.sheet_properties.tabColor = NAVY
        headers = ["Message"] + FIELD_HEADERS
        set_widths(ws, [20] + FIELD_WIDTHS)
        row = banner(ws, "All fields", "Every field of every message · use the filter buttons to search",
                     len(headers), Link("← Overview", "Overview"))
        rows, bands, encs = [], [], []
        for b, frame in enumerate(self.p.frames):
            for i, (f, offset, size) in enumerate(frame_offsets(frame)):
                rows.append([Link(frame.name, self.frame_sheet[frame.name])] + field_row(self.p, frame, i, offset, size))
                bands.append(b)
                encs.append(f.encoding)

        def style(r, c, value):
            if c == 3:
                return {"fill": fill(ENCODING_FILL[encs[r]]), "font": font(bold=True, color=ENCODING_INK[encs[r]])}
            if c == 8:
                return {"font": font(color=ENCODING_INK[encs[r]])}
            if c in (2, 4, 5):
                return {"alignment": Alignment(horizontal="center", vertical="center")}
            return {}

        table(ws, row, headers, rows, style=style, band=bands, filter_=True, freeze=True)

    # ---------------------------------------------------------- reference ---

    def enums_sheet(self) -> None:
        ws = self.wb.create_sheet("Enumerations")
        ws.sheet_properties.tabColor = "8250DF"
        widths = [22, 12, 12, 26, 40, 40]
        set_widths(ws, widths)
        row = banner(ws, "Enumerations", "Numbers on the wire and the names shown instead", len(widths),
                     Link("← Overview", "Overview"))
        if not self.p.enums:
            write(ws, row, 1, "This protocol has no enumerations.", font=font(color=MUTED))
            return
        rows, bands = [], []
        for b, (name, enum) in enumerate(self.p.enums.items()):
            used = ", ".join(f"{fr.name}.{fd.name}" for fr in self.p.frames for fd in fr.fields
                             if fd.enum == name or any(bit.enum == name for bit in fd.bits))
            for k, (value, label) in enumerate(sorted(enum.values.items())):
                rows.append([name if k == 0 else "", _hex(value), value, label,
                             enum.description if k == 0 else "", used if k == 0 else ""])
                bands.append(b)

        def style(r, c, value):
            if c == 1 and value:
                return {"font": font(bold=True, color="8250DF")}
            if c in (2, 3):
                return {"alignment": Alignment(horizontal="center", vertical="center")}
            if c == 4:
                return {"font": font(bold=True)}
            return {}

        table(ws, row, ["Enumeration", "Value (hex)", "Value (dec)", "Name", "Description", "Used by"], rows,
              style=style, band=bands, filter_=True, freeze=True)

    def crc_sheet(self) -> None:
        ws = self.wb.create_sheet("CRC")
        ws.sheet_properties.tabColor = ENCODING_INK[Encoding.CRC]
        widths = [18, 8, 14, 14, 9, 9, 14, 14, 50]
        set_widths(ws, widths)
        row = banner(ws, "Checksum / CRC algorithms", 'Parameters (Rocksoft model). "Check" is the result for the '
                     'ASCII text "123456789" – use it to verify your implementation.', len(widths),
                     Link("← Overview", "Overview"))
        algs: Dict[str, Any] = {}
        for frame in self.p.frames:
            for f in frame.fields:
                if f.encoding == Encoding.CRC:
                    try:
                        alg = crc_algorithm(f)
                    except DefinitionError:
                        continue
                    algs.setdefault(alg.name, [alg, []])[1].append(f"{frame.name}.{f.name}")
        if not algs:
            write(ws, row, 1, "This protocol has no checksum fields.", font=font(color=MUTED))
            return
        rows = []
        for name, (alg, users) in sorted(algs.items()):
            w = (alg.width + 3) // 4
            check = "" if alg.check is None else f"0x{alg.check:0{alg.size * 2}X}"
            if alg.custom is None:
                rows.append([name, alg.width, f"0x{alg.poly:0{w}X}", f"0x{alg.init:0{w}X}",
                             "yes" if alg.refin else "no", "yes" if alg.refout else "no",
                             f"0x{alg.xorout:0{w}X}", check, ", ".join(users)])
            else:
                rows.append([name, alg.width, alg.description, "", "", "", "", check, ", ".join(users)])
        table(ws, row, ["Algorithm", "Width (bits)", "Polynomial", "Init", "Reflect in", "Reflect out",
                        "XOR out", "Check", "Used by"], rows,
              style=lambda r, c, v: {"font": font(bold=True, color=ENCODING_INK[Encoding.CRC])} if c == 1 else
              {"font": font(name=MONO)} if c in (3, 4, 7, 8) else
              {"alignment": Alignment(horizontal="center", vertical="center")} if c in (2, 5, 6) else {})

    def tests_sheet(self) -> None:
        ws = self.wb.create_sheet("Tests")
        ws.sheet_properties.tabColor = "9A6700"
        widths = [22, 7, 12, 20, 44, 44, 11]
        set_widths(ws, widths)
        row = banner(ws, "Tests", "Automated test sequences (run them on the Test page or with 'pdcli test')",
                     len(widths), Link("← Overview", "Overview"))
        if not self.p.tests:
            write(ws, row, 1, "This protocol has no tests.", font=font(color=MUTED))
            return
        rows, bands = [], []
        for b, test in enumerate(self.p.tests):
            for i, step in enumerate(test.steps, 1):
                d = step_to_dict(step)
                action = d.pop("action")
                message = d.pop("message", "")
                checks = d.pop("checks", [])
                timeout = d.pop("timeout_ms", test.timeout_ms if action in ("expect", "transact") else "")
                rows.append([test.name if i == 1 else "", i, action, message, _step_details(d),
                             "\n".join(_check_text(c) for c in checks), f"{timeout} ms" if timeout != "" else ""])
                bands.append(b)

        def style(r, c, value):
            if c == 1 and value:
                return {"font": font(bold=True, color="9A6700")}
            if c in (2, 7):
                return {"alignment": Alignment(horizontal="center", vertical="center")}
            if c == 3:
                return {"font": font(bold=True)}
            return {}

        table(ws, row, ["Test", "Step", "Action", "Message", "Values / details", "Checks", "Timeout"], rows,
              style=style, band=bands, filter_=True, freeze=True)

    def simulator_sheet(self) -> None:
        ws = self.wb.create_sheet("Simulator")
        ws.sheet_properties.tabColor = "6E7781"
        widths = [20, 26, 20, 20, 60, 10]
        set_widths(ws, widths)
        row = banner(ws, "Device simulator", "How the built-in simulated device answers (first matching rule wins)",
                     len(widths), Link("← Overview", "Overview"))
        if not self.p.simulator:
            write(ws, row, 1, "This protocol has no simulator rules.", font=font(color=MUTED))
            return
        rows = [[r.on, "\n".join(f"{k} = {_value_text(v)}" for k, v in r.match.items()), r.reply or "(no reply)",
                 ", ".join(r.copy), "\n".join(f"{k} = {_value_text(v)}" for k, v in r.values.items()),
                 f"{r.delay_ms} ms" if r.delay_ms else ""] for r in self.p.simulator]
        table(ws, row, ["On request", "Only if", "Reply", "Copied fields", "Reply values", "Delay"], rows,
              style=lambda r, c, v: {"font": font(bold=True)} if c in (1, 3) else {})

    def validation_sheet(self) -> None:
        ws = self.wb.create_sheet("Validation")
        errors = sum(1 for i in self.issues if i.severity == ERROR)
        ws.sheet_properties.tabColor = "B91C1C" if errors else "15803D"
        widths = [12, 34, 90]
        set_widths(ws, widths)
        row = banner(ws, "Validation", f"{errors} error(s), {len(self.issues) - errors} warning(s)", len(widths),
                     Link("← Overview", "Overview"))
        if not self.issues:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            write(ws, row, 1, "✓  No problems found", font=font(bold=True, size=14, color="15803D"),
                  fill=fill("E6F4EA"), alignment=Alignment(vertical="center", indent=1))
            ws.row_dimensions[row].height = 30
            return
        rows = [[i.severity, i.location, i.message] for i in self.issues]
        table(ws, row, ["Severity", "Location", "Problem"], rows, filter_=True, freeze=True,
              style=lambda r, c, v: ({"font": font(bold=True, color="B91C1C" if v == ERROR else "9A6700"),
                                      "fill": fill("FDE2E1" if v == ERROR else "FEF3C7")} if c == 1 else {}))


def build_workbook(protocol: ProtocolDefinition, source: Optional[str] = None) -> Workbook:
    return WorkbookBuilder(protocol, source).build()


def export_xlsx(protocol: ProtocolDefinition, path: Union[str, Path], source: Optional[str] = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(protocol, source).save(str(path))
    return path
