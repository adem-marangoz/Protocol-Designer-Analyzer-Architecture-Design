"""Protocol export to a formatted Excel workbook (.xlsx)."""

import pytest
from openpyxl import load_workbook

from helpers import make_protocol

from protocol_designer.application.excel_export import (
    ENCODING_FILL, FIXED_SHEETS, SheetNames, build_workbook, example_packet, export_xlsx,
)
from protocol_designer.application.spec_report import export
from protocol_designer.protocol.model import Encoding
from protocol_designer.storage import load_protocol

EXAMPLES = ["tpms_rs485", "tpms_can", "ecu_bootloader"]


@pytest.fixture
def book(tmp_path, tpms):
    """The RS485 example, written to disk and read back like Excel would."""
    return load_workbook(export_xlsx(tpms, tmp_path / "tpms.xlsx", "tpms_rs485.json"))


def values(ws):
    return [c.value for row in ws.iter_rows() for c in row if c.value is not None]


def text(ws):
    return "\n".join(str(v) for v in values(ws))


def find(ws, value):
    for row in ws.iter_rows():
        for c in row:
            if c.value == value:
                return c
    raise AssertionError(f"{value!r} not on sheet {ws.title}")


def table(ws, first_header):
    """The table whose header row starts with ``first_header``, as a list of {header: value} dicts."""
    head = find(ws, first_header)
    headers = [c.value for c in ws[head.row][head.column - 1:] if c.value is not None]
    rows = []
    for r in range(head.row + 1, ws.max_row + 1):
        cells = [ws.cell(row=r, column=head.column + i).value for i in range(len(headers))]
        if all(v is None for v in cells):
            break
        rows.append(dict(zip(headers, cells)))
    return rows


def check_links(wb):
    """Every internal hyperlink must point to an existing sheet; returns the targets."""
    targets = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink is None:
                    continue
                link = c.hyperlink.location  # internal link, not a web address
                assert link and not c.hyperlink.target, link
                sheet, cell = link.rsplit("!", 1)
                sheet = sheet[1:-1].replace("''", "'")
                assert sheet in wb.sheetnames, (ws.title, c.coordinate, link)
                assert wb[sheet][cell] is not None
                targets.append((ws.title, sheet))
    return targets


def test_sheets_in_reading_order(book, tpms):
    assert book.sheetnames == (["Overview", "Messages", "Byte maps"] + [f.name for f in tpms.frames]
                               + ["All fields", "Enumerations", "CRC", "Tests", "Simulator", "Validation"])
    assert book.active.title == "Overview"
    for ws in book.worksheets:
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.fitToWidth == 1 and ws.sheet_properties.pageSetUpPr.fitToPage
        assert ws["A1"].value  # navy title band
    assert book.properties.title.startswith("TPMS RS485 Protocol")


def test_overview_tiles_contents_and_legend(book, tpms):
    ws = book["Overview"]
    assert ws["A1"].value == "TPMS RS485 Protocol"
    tiles = {ws.cell(row=5, column=c).value: ws.cell(row=4, column=c).value for c in range(2, 7)}
    assert tiles == {"Messages": len(tpms.frames), "Fields": sum(len(f.fields) for f in tpms.frames),
                     "Enumerations": len(tpms.enums), "Tests": len(tpms.tests), "Problems": 0}
    content = text(ws)
    for needle in ["Example tyre-pressure sensor", "Big endian", "tpms_rs485.json", "SIMULATOR",
                   "115200 baud", "Contents (click to open)", "Colour legend", "How to read this workbook"]:
        assert needle in content, needle
    # the contents list links to every other sheet
    linked = {target for source, target in check_links(book) if source == "Overview"}
    assert linked == set(book.sheetnames) - {"Overview"}


def test_every_link_works_and_sheets_link_back(book):
    targets = check_links(book)
    for name in book.sheetnames[1:]:
        assert (name, "Overview") in targets or (name, "Messages") in targets, name


def test_messages_table(book, tpms):
    rows = table(book["Messages"], "Message")
    assert [r["Message"] for r in rows] == [f.name for f in tpms.frames]
    first = rows[0]
    assert first["ID (hex)"] == "0x10" and first["ID (dec)"] == 16 and first["Direction"] == "TX"
    assert first["Size"] == "6 bytes" and first["Expected response"] == "SENSOR_DATA"
    assert book["Messages"].auto_filter.ref and book["Messages"].freeze_panes
    cell = find(book["Messages"], "READ_SENSOR")
    assert cell.hyperlink.location == "'READ_SENSOR'!A1"
    assert any(r["Size"] == "11 bytes + variable" for r in rows)


def test_message_sheet_fields_bits_and_example(book, tpms):
    ws = book["SENSOR_DATA"]
    assert ws.sheet_properties.tabColor.rgb.endswith("15803D")  # RX = green
    content = text(ws)
    frame = tpms.frame("SENSOR_DATA")
    encoded, _ = example_packet(tpms, frame)
    assert encoded.hex in content  # example packet
    assert "RX – received from the device" in content and "READ_SENSOR" in content  # response to
    rows = {r["Field"]: r for r in table(ws, "#")}
    pressure = rows["PRESSURE"]
    assert (pressure["Byte offset"], pressure["Size (bytes)"], pressure["Type"]) == (5, 2, "UINT16 BE")
    assert pressure["Byte order"] == "big endian" and pressure["Encoding"] == "Value"
    assert pressure["Scale"] == 0.1 and pressure["Unit"] == "bar" and (pressure["Min"], pressure["Max"]) == (0, 50)
    assert "value = raw × 0.1" in pressure["Rule / value"]
    assert pressure["Example (hex)"] == "00 00" and pressure["Example value"] == "0.0 bar"
    assert rows["CRC"]["Encoding"] == "CRC" and "CRC8 over SOF" in rows["CRC"]["Rule / value"]
    assert rows["STATUS"]["Enumeration"] == "SENSOR_STATUS"
    # bits of FLAGS follow it as indented rows
    state = rows["↳ Sensor State"]
    assert state["Size (bytes)"] == "bits 3–4" and state["Byte offset"] == 10
    assert "mask 0x18" in state["Rule / value"] and "0 = SLEEP" in state["Rule / value"]
    assert state["Enumeration"] == "SENSOR_STATE"
    assert rows["↳ Sensor Active"]["Type"] == "flag"
    # rules
    assert "CRC = CRC8 over the bytes of SOF" in content and "LENGTH = number of bytes of SENSOR_ID" in content
    # colours follow the encoding
    sof = find(ws, "SOF")
    assert sof.fill.fgColor.rgb.endswith(ENCODING_FILL[Encoding.CONSTANT])


def test_byte_map_merges_fields_and_shows_example(book, tpms):
    ws = book["Byte maps"]
    frame = tpms.frame("SENSOR_DATA")
    title = find(ws, "SENSOR_DATA   ·   RX   ·   12 bytes")
    assert title.hyperlink.location == "'SENSOR_DATA'!A1"
    byte_row, field_row, example_row = title.row + 1, title.row + 2, title.row + 3
    assert [ws.cell(row=byte_row, column=c).value for c in range(1, 14)] == ["Byte"] + list(range(12))
    encoded, _ = example_packet(tpms, frame)
    assert [ws.cell(row=example_row, column=c).value for c in range(2, 14)] == [f"{b:02X}" for b in encoded.data]
    pressure = find(ws, "PRESSURE")
    assert pressure.row == field_row
    merged = [str(r) for r in ws.merged_cells.ranges if r.min_row == field_row]
    assert f"{pressure.coordinate}:{ws.cell(row=field_row, column=pressure.column + 1).coordinate}" in merged
    # the message sheet links back to its byte map
    link = find(book["SENSOR_DATA"], "Byte map →").hyperlink.location
    assert link == f"'Byte maps'!A{title.row}"


def test_long_frame_wraps_and_marks_continuation(tmp_path):
    p = make_protocol([{"name": "BIG", "fields": [{"name": "HEAD", "type": "UINT8", "encoding": "CONSTANT",
                                                    "value": "0x55"},
                                                   {"name": "BLOB", "type": "BYTES", "size": 20}]}])
    ws = load_workbook(export_xlsx(p, tmp_path / "big.xlsx"))["Byte maps"]
    assert find(ws, "BLOB") and find(ws, "BLOB (cont.)")
    assert find(ws, 16).row == find(ws, 0).row + 3  # second line of 16 bytes


def test_all_fields_enums_crc_tests_simulator(book, tpms):
    rows = table(book["All fields"], "Message")
    assert len(rows) == sum(len(f.fields) for f in tpms.frames)
    assert {r["Message"] for r in rows} == {f.name for f in tpms.frames}
    assert book["All fields"].auto_filter.ref
    enums = text(book["Enumerations"])
    assert "SENSOR_STATE" in enums and "PARKED" in enums and "SENSOR_DATA.FLAGS" in enums
    crc = table(book["CRC"], "Algorithm")
    assert crc[0]["Algorithm"] == "CRC8" and crc[0]["Check"] == "0xF4" and crc[0]["Polynomial"] == "0x07"
    assert crc[0]["Reflect in"] == "no"
    tests = text(book["Tests"])
    assert "Sensor Read" in tests and "PRESSURE between 1.5 and 3.5" in tests and "STATUS = ACTIVE" in tests
    sim = text(book["Simulator"])
    assert "FLAGS = (Sensor Active = true" in sim and "TEMPERATURE = 25.0" in sim
    assert "No problems found" in text(book["Validation"])


def test_can_and_little_endian_details(tmp_path, protocols_dir):
    wb = load_workbook(export_xlsx(load_protocol(protocols_dir / "tpms_can.json"), tmp_path / "can.xlsx"))
    rows = {r["Message"]: r for r in table(wb["Messages"], "Message")}
    assert rows["WHEEL_PRESSURE"]["CAN ID"] == "0x300"
    assert rows["VEHICLE_SPEED"]["CAN ID"] == "0x19000000 (ext)"
    assert "0x300 (11-bit standard)" in text(wb["WHEEL_PRESSURE"])
    assert "little endian" in text(wb["All fields"])
    assert "recognised by its CAN identifier" in text(wb["VEHICLE_SPEED"])


def test_bootloader_sequence_reads_naturally(tmp_path, protocols_dir):
    wb = load_workbook(export_xlsx(load_protocol(protocols_dir / "ecu_bootloader.json"), tmp_path / "b.xlsx"))
    tests = text(wb["Tests"])
    assert "key = ${seed} ^ 0x5A5A5A5A" in tests  # set step
    assert "SEED ≠ 0 → save as ${seed}" in tests  # check with save_as
    assert "Flashing finished" in tests  # log step
    assert "500 ms" in tests  # the test's default timeout for transact steps
    assert "CRC16_MODBUS" in text(wb["CRC"]) and "0x4B37" in text(wb["CRC"])


def test_sheet_names_are_valid_and_unique():
    names = SheetNames(FIXED_SHEETS)
    assert names.make("A/B:C*D?") == "A_B_C_D_"
    assert names.make("x" * 40) == "x" * 31
    assert names.make("X" * 40) == "X" * 27 + " (2)"  # unique ignoring case, still 31 characters
    assert names.make("messages") == "messages (2)"  # clashes with a fixed sheet, ignoring case
    assert names.make("History") == "History (2)"  # reserved by Excel
    assert names.make("'quoted'") == "quoted"


def test_awkward_message_names(tmp_path):
    frames = [{"name": n, "fields": [{"name": "A", "type": "UINT8"}]}
              for n in ["Overview", "VERY_LONG_MESSAGE_NAME_THAT_DOES_NOT_FIT", "VERY_LONG_MESSAGE_NAME_THAT_DOES_NOT_FIT_2"]]
    wb = load_workbook(export_xlsx(make_protocol(frames), tmp_path / "names.xlsx"))
    assert "Overview (2)" in wb.sheetnames
    assert all(len(n) <= 31 for n in wb.sheetnames) and len(set(n.lower() for n in wb.sheetnames)) == len(wb.sheetnames)
    check_links(wb)
    assert find(wb["Overview (2)"], "Overview")  # its title is the message name


def test_text_is_never_a_formula_and_control_characters_are_removed(tmp_path):
    p = make_protocol([{"name": "M", "description": "=HYPERLINK(\"x\")",
                        "fields": [{"name": "A", "type": "UINT8", "description": "bad\x01char"}]}])
    ws = load_workbook(export_xlsx(p, tmp_path / "f.xlsx"))["M"]
    cell = find(ws, "=HYPERLINK(\"x\")")
    assert cell.data_type == "s"
    assert any(isinstance(v, str) and v.startswith("bad") and "\x01" not in v for v in values(ws))


def test_invalid_protocol_still_exports_with_its_problems(tmp_path):
    p = make_protocol([{"name": "M", "fields": [{"name": "X", "type": "UINT8"}, {"name": "X", "type": "UINT8"},
                                                 {"name": "C", "type": "UINT8", "encoding": "CONSTANT",
                                                  "value": "0x1FF"}]}])
    wb = load_workbook(export_xlsx(p, tmp_path / "bad.xlsx"))
    assert "duplicate field name" in text(wb["Validation"])
    assert "none: the message definition has errors" in text(wb["M"])
    problems = wb["Overview"]["F4"]
    assert problems.value > 0 and problems.font.color.rgb.endswith("B91C1C")
    assert wb["Validation"].sheet_properties.tabColor.rgb.endswith("B91C1C")


def test_empty_protocol(tmp_path):
    wb = load_workbook(export_xlsx(make_protocol([]), tmp_path / "empty.xlsx"))
    assert "This protocol has no tests." in text(wb["Tests"])
    assert "This protocol has no checksum fields." in text(wb["CRC"])


@pytest.mark.parametrize("name", EXAMPLES)
def test_examples_export(tmp_path, protocols_dir, name):
    protocol = load_protocol(protocols_dir / f"{name}.json")
    path = export(protocol, tmp_path / f"{name}.xlsx", f"{name}.json")  # chosen by the extension
    assert path.read_bytes()[:2] == b"PK"
    wb = load_workbook(path)
    for frame in protocol.frames:
        assert frame.name in wb.sheetnames
    check_links(wb)


def test_build_workbook_in_memory(tpms):
    wb = build_workbook(tpms)
    assert wb.sheetnames[0] == "Overview" and len(wb.sheetnames) == len(FIXED_SHEETS) + len(tpms.frames)
